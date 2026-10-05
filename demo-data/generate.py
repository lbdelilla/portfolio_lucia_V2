"""Generates the fictitious dataset behind the Looker Studio demo dashboard.

Every number here is invented. The structure imitates the operation of a
training academy (cohorts, students, mentoring, certificates) so the dashboard
can show realistic analysis without using any real company data.

    python demo-data/generate.py
"""
import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)

OUT = Path(__file__).parent / 'academia_demo.xlsx'
PROGRAMS = {
    'Desarrollo Full Stack': 18,
    'Ciencia de Datos': 16,
    'Ingeniería de IA': 14,
    'Ciberseguridad': 16,
}
REGIONS = ['España', 'Latinoamérica']
MODES = ['Part-time', 'Full-time']
# From this month on, certificate reminders are automated in the story the data tells
AUTOMATION = date(2024, 9, 1)
FIRST = date(2024, 1, 1)
LAST = date(2026, 6, 30)


def month_starts(first, last):
    current = date(first.year, first.month, 1)
    while current <= last:
        yield current
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)


def clamp(value, low, high):
    return max(low, min(high, value))


# ---------- Cohorts: one row per cohort
cohorts = []
number = 1
for start in month_starts(FIRST, date(2026, 3, 1)):
    for _ in range(random.choice([2, 2, 3, 3, 4])):
        program = random.choice(list(PROGRAMS))
        region = random.choices(REGIONS, weights=[6, 4])[0]
        mode = random.choices(MODES, weights=[7, 3])[0]
        weeks = PROGRAMS[program] if mode == 'Part-time' else PROGRAMS[program] // 2
        begins = start + timedelta(days=random.randint(0, 20))
        ends = begins + timedelta(weeks=weeks)
        enrolled = random.randint(12, 32)
        dropouts = round(enrolled * clamp(random.gauss(0.09, 0.04), 0.0, 0.25))
        finished = ends <= LAST
        graduates = round((enrolled - dropouts) * clamp(random.gauss(0.88, 0.06), 0.6, 1.0)) if finished else None
        # Ratings improve slowly over time; teachers score a little higher than the overall experience
        progress = (begins - FIRST).days / (LAST - FIRST).days
        rating = round(clamp(random.gauss(8.3 + 0.5 * progress, 0.3), 6.5, 10), 2)
        teacher_rating = round(clamp(rating + random.gauss(0.45, 0.25), 6.5, 10), 2)
        if not finished:
            days_to_certificate = None
        elif ends < AUTOMATION:
            days_to_certificate = round(clamp(random.gauss(21, 5), 9, 40), 1)
        else:
            days_to_certificate = round(clamp(random.gauss(6, 2.2), 1, 14), 1)
        cohorts.append(
            {
                'cohorte_id': f'C-{number:03d}',
                'programa': program,
                'region': region,
                'modalidad': mode,
                'fecha_inicio': begins,
                'fecha_fin': ends,
                'estado': 'Finalizada' if finished else 'En curso',
                'inscritos': enrolled,
                'bajas': dropouts,
                'graduados': graduates,
                'valoracion_general': rating,
                'valoracion_profesores': teacher_rating,
                'dias_hasta_certificado': days_to_certificate,
                'dias_antelacion_mentor_confirmado': clamp(round(random.gauss(12 + 6 * progress, 5)), 0, 35),
            }
        )
        number += 1

cohorts_df = pd.DataFrame(cohorts)

# ---------- Monthly: one row per month, region and programme
monthly = []
for month in month_starts(FIRST, LAST):
    month_end = date(month.year + (month.month == 12), month.month % 12 + 1, 1) - timedelta(days=1)
    for region in REGIONS:
        for program in PROGRAMS:
            running = cohorts_df[
                (cohorts_df.region == region)
                & (cohorts_df.programa == program)
                & (cohorts_df.fecha_inicio <= month_end)
                & (cohorts_df.fecha_fin >= month)
            ]
            active = int((running.inscritos - running.bajas).sum())
            if active == 0:
                continue
            offered = round(active * random.uniform(0.9, 1.5))
            done = round(offered * clamp(random.gauss(0.93, 0.04), 0.75, 1.0))
            incidents = round(active * random.uniform(0.04, 0.12))
            # Incidents get resolved faster after the automation date
            fast_share = 0.62 if month < AUTOMATION else 0.86
            fast = round(incidents * clamp(random.gauss(fast_share, 0.07), 0.3, 1.0))
            monthly.append(
                {
                    'mes': month,
                    'region': region,
                    'programa': program,
                    'cohortes_activas': len(running),
                    'estudiantes_activos': active,
                    'mentorias_ofrecidas': offered,
                    'mentorias_realizadas': done,
                    'horas_mentoria': round(done * random.uniform(0.7, 0.95), 1),
                    'incidencias': incidents,
                    'incidencias_resueltas_48h': fast,
                }
            )

monthly_df = pd.DataFrame(monthly)

# ---------- Notes sheet
notes = pd.DataFrame(
    [
        ['Qué es esto', 'Datos de ejemplo para un tablero de demostración en Looker Studio.'],
        ['Aviso', 'Todos los datos son ficticios. No proceden de ninguna empresa real ni representan a personas reales.'],
        ['Hoja "cohortes"', 'Una fila por cohorte: programa, región, fechas, inscritos, bajas, graduados, valoraciones y días hasta el certificado.'],
        ['Hoja "mensual"', 'Una fila por mes, región y programa: estudiantes activos, mentorías e incidencias.'],
        ['Valoraciones', 'Escala de 0 a 10.'],
        ['dias_hasta_certificado', 'Días entre el fin de la cohorte y la emisión del certificado. Vacío si la cohorte sigue en curso.'],
        ['Supuesto del ejemplo', 'A partir de septiembre de 2024 se simula la automatización de avisos y recordatorios: bajan los días hasta el certificado y sube la proporción de incidencias resueltas en 48 horas.'],
        ['Cómo se generó', 'Con el script demo-data/generate.py del repositorio del portfolio, con una semilla fija para que el resultado sea reproducible.'],
    ],
    columns=['Campo', 'Descripción'],
)

with pd.ExcelWriter(OUT, engine='openpyxl', date_format='yyyy-mm-dd', datetime_format='yyyy-mm-dd') as writer:
    notes.to_excel(writer, sheet_name='leeme', index=False)
    cohorts_df.to_excel(writer, sheet_name='cohortes', index=False)
    monthly_df.to_excel(writer, sheet_name='mensual', index=False)

# ---------- Formatting: Arial throughout, clear headers, readable widths
book = load_workbook(OUT)
for sheet in book.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name='Arial', size=10)
    for cell in sheet[1]:
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='221C3A')
        cell.alignment = Alignment(vertical='center')
    sheet.freeze_panes = 'A2'
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 90 if sheet.title == 'leeme' else 34)
    if sheet.title == 'leeme':
        for row in sheet.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical='top')
book.save(OUT)

print(f'{OUT.name}: {len(cohorts_df)} cohortes, {len(monthly_df)} filas mensuales')
